"""Inventory the supplied editorial and design corpus for release audits.

Writes a readable manifest plus extracted DOCX text so route coverage and missing
source material can be checked without opening each binary by hand.
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CORPUS = (
    ROOT.parent
    / "EU4Youth"
    / "site web EU4Youth.org-20260811T205332Z-1-001"
    / "site web EU4Youth.org"
)
DESIGNS = (
    ROOT.parent
    / "EU4Youth"
    / "UI Web Design-20260807T094637Z-1-001"
    / "UI Web Design"
)
OUT = ROOT / "tools" / "out" / "source-audit"


def relative(path: Path, base: Path) -> str:
    return path.relative_to(base).as_posix()


def docx_text(path: Path) -> str:
    document = Document(path)
    lines = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        lines.append("[TABLE]")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(line for line in lines if line)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    manifest: list[str] = ["EU4YOUTH SUPPLIED-SOURCE INVENTORY", ""]

    manifest.append("DESIGN PDFS")
    for path in sorted(DESIGNS.glob("*.pdf")):
        document = pymupdf.open(path)
        page_sizes = ", ".join(
            f"{round(page.rect.width)}x{round(page.rect.height)}" for page in document
        )
        manifest.append(
            f"- {path.name}: {len(document)} page(s), {page_sizes}, "
            f"{path.stat().st_size / 1024 / 1024:.1f} MB"
        )
        document.close()

    manifest.extend(["", "EDITORIAL DOCX"])
    for path in sorted(CORPUS.rglob("*.docx")):
        text = docx_text(path)
        target = OUT / f"{relative(path, CORPUS).replace('/', '__')}.txt"
        target.write_text(text, encoding="utf8")
        words = len(text.split())
        manifest.append(f"- {relative(path, CORPUS)}: {words} extracted words")

    manifest.extend(["", "SOURCE PDFS"])
    for path in sorted(CORPUS.rglob("*.pdf")):
        document = pymupdf.open(path)
        manifest.append(
            f"- {relative(path, CORPUS)}: {len(document)} page(s), "
            f"{path.stat().st_size / 1024 / 1024:.1f} MB"
        )
        document.close()

    manifest.extend(["", "WORKBOOKS"])
    for path in sorted(CORPUS.rglob("*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=True)
        manifest.append(f"- {relative(path, CORPUS)}")
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            nonempty = [row for row in rows if any(value is not None for value in row)]
            header_candidates = nonempty[:20]
            first_row = (
                max(header_candidates, key=lambda row: sum(value is not None for value in row))
                if header_candidates
                else ()
            )
            headers = [
                str(value).strip()
                for value in first_row
                if value is not None
            ]
            columns = max((len(row) for row in nonempty), default=0)
            manifest.append(
                f"  - {sheet.title}: {len(nonempty)} non-empty rows x {columns} cols; "
                f"headers: {', '.join(headers[:12])}"
            )
        workbook.close()

    html_files = sorted(CORPUS.rglob("*.html"))
    if html_files:
        manifest.extend(["", "HTML"])
        for path in html_files:
            manifest.append(
                f"- {relative(path, CORPUS)}: {path.stat().st_size / 1024:.0f} KB"
            )

    target = OUT / "manifest.txt"
    target.write_text("\n".join(manifest), encoding="utf8")
    print(f"{target}: {len(manifest)} manifest lines")


if __name__ == "__main__":
    main()
